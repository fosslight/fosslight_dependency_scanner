#!/usr/bin/env python
# -*- coding: utf-8 -*-
# Copyright (c) 2026 LG Electronics Inc.
# SPDX-License-Identifier: Apache-2.0
"""Assertions for fosslight_dependency Excel reports."""

from pathlib import Path

import openpyxl

DEP_SHEET_NAME = "DEP_FL_Dependency"


def assert_dep_sheet_has_min_rows(
    output_path: str,
    min_rows: int = 2,
    expect_excel: bool = True,
) -> None:
    """Fail if the newest dep Excel report has fewer than ``min_rows`` on DEP_FL_Dependency.

    When ``expect_excel`` is True (default), a missing ``fosslight_report_dep_*.xlsx``
    fails the assertion. Pass ``expect_excel=False`` for non-Excel formats such as
    ``-f opossum``.
    """
    reports = list(Path(output_path).glob("fosslight_report_dep_*.xlsx"))
    if not reports:
        assert not expect_excel, (
            f"No fosslight_report_dep_*.xlsx found under {output_path}"
        )
        return

    report = max(reports, key=lambda path: path.stat().st_mtime)
    workbook = openpyxl.load_workbook(report, read_only=True, data_only=True)
    try:
        assert DEP_SHEET_NAME in workbook.sheetnames, (
            f"{DEP_SHEET_NAME} sheet not found in {report}; "
            f"sheets={workbook.sheetnames}"
        )
        sheet = workbook[DEP_SHEET_NAME]
        row_count = 0
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None and str(cell).strip() != "" for cell in row):
                row_count += 1
        assert row_count >= min_rows, (
            f"{DEP_SHEET_NAME} in {report} has {row_count} row(s); "
            f"expected at least {min_rows} (header + data)"
        )
    finally:
        workbook.close()
